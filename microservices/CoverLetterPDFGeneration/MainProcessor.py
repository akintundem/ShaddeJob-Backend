from docx import Document
import pypandoc

def convert_to_pdf(input_path, output_path):
    try:
        pypandoc.convert_file(input_path, 'pdf', outputfile=output_path)
        print("PDF conversion successful.")
    except Exception as e:
        print(f"Error: {e}")


def populate_docx(template_path, output_path, data):
    doc = Document(template_path)
    
    # Replace placeholders
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            placeholder = f'{{{{{key}}}}}'
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, value)

    doc.save(output_path)



data = {
    "Name": "Mayokun Moses Akintunde",
    "phone": "(204) 869-1025",
    "email": "akintundemayokun@gmail.com",
    "CompanyName": "Manitoba Hydro",
    "Body": ("I am writing to express my interest in the position of System Developer at Manitoba Hydro. "
             "Having recently completed my bachelor’s degree in computer engineering at the University of Manitoba, "
             "I bring a solid foundation in computer science principles and hands-on engineering experience honed through various academic projects, "
             "personal projects, and professional engagements, including my previous internship at Manitoba Hydro and Greif.\n\n"
             "In my previous role at Manitoba Hydro, I was tasked with developing a solution for managing generator study data for the Integrated Resource Planning Division. "
             "I designed and implemented a specialized web application as an internal tool for engineers, using a Microsoft SQL Server database for storage and a Flask-based web server responsible for user authentication, "
             "authorization, report generation, and the creation of specialized files for seamless integration into the PSS\E software. "
             "I documented technical processes and API specifications while collaborating with engineers from various departments to comprehend their requirements and facilitate the creation of a user-friendly web-based client tailored to meet their needs. "
             "This experience has equipped me with a deep understanding of system analysis, design, development, testing, implementation, and documentation.\n\n"
             "I am proficient in SQL and have extensive experience working with RDBMS such as Microsoft SQL Server. "
             "My technical skill set includes Python, Java, C/C++, and experience with Flask which is an alternative to Django will be beneficial for developing and maintaining the data-related systems at your company. "
             "Additionally, I have developed and utilized REST APIs, worked with version control systems like Git, and have hands-on experience with cloud service providers including AWS. "
             "I have hands-on experience completing projects using Agile methodologies, ensuring high-quality software delivery. As a Scrum Leader and Lead Programmer for an Airline Reservation System, "
             "I managed a team and oversaw the development of a comprehensive Android app, ensuring the project met all client requirements.\n\n"
             "During my time at Manitoba Hydro, I worked closely with top engineers whose mentorship greatly advanced my technical and professional skills. "
             "The collaborative environment enabled me to build strong relationships, and I thoroughly enjoyed the challenging work and the company's strategic position in the global energy landscape. "
             "As energy demand surges due to AI and population growth, Manitoba Hydro stands at the forefront of this critical industry. Joining your team would deepen my understanding of the energy sector and help me develop the skills to grow within the company while contributing to its future success.\n\n"
             "I am enthusiastic about the opportunity to return to Manitoba Hydro as a System Developer and am confident that my skills and experiences make me a strong candidate for this position. "
             "I look forward to the possibility of discussing how my skills will be beneficial to your team over an interview.")
}

template_path = 'Templates/template.docx'
docx_output_path = 'cover_letter_filled.docx'
pdf_output_path = 'cover_letter_filled.pdf'

populate_docx(template_path, docx_output_path, data)
convert_to_pdf(docx_output_path, pdf_output_path)
